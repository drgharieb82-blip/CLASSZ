import csv
import io
import json
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.db import migrations  # noqa: F401
from app.modules.chapters.models import Chapter
from app.modules.concepts.models import Concept
from app.modules.courses.models import Course
from app.modules.lessons.models import Lesson
from app.modules.question_bank.models import Difficulty, QuestionCategory, QuestionType
from app.modules.question_bank.schemas import QuestionChoiceCreate, QuestionCreate
from app.modules.question_bank.service import create_question
from app.modules.question_import.models import ImportError, ImportJob, ImportJobStatus, ImportSourceType, ImportSummary


def _job_options():
    return (
        selectinload(ImportJob.errors),
        selectinload(ImportJob.summary),
    )


def _source_type(filename: str) -> ImportSourceType:
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix == "json":
        return ImportSourceType.JSON
    if suffix in {"xlsx", "xls"}:
        return ImportSourceType.EXCEL
    return ImportSourceType.CSV


def _parse_rows(filename: str, content: bytes) -> list[dict]:
    source_type = _source_type(filename)
    if source_type == ImportSourceType.JSON:
        parsed = json.loads(content.decode("utf-8-sig"))
        if isinstance(parsed, dict):
            parsed = parsed.get("questions", [])
        return [dict(row) for row in parsed]

    if source_type == ImportSourceType.EXCEL:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - depends on optional local package
            raise ValueError("Excel imports require openpyxl to be installed") from exc

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
            if any(value is not None for value in row)
        ]

    stream = io.StringIO(content.decode("utf-8-sig"))
    return [dict(row) for row in csv.DictReader(stream)]


def _get(row: dict, *names: str) -> str | None:
    lowered = {str(key).strip().lower(): value for key, value in row.items()}
    for name in names:
        value = lowered.get(name.lower())
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def _split(value: str | None) -> list[str]:
    if not value:
        return []
    if value.strip().startswith("["):
        parsed = json.loads(value)
        return [str(item).strip() for item in parsed if str(item).strip()]
    return [part.strip() for part in value.replace("|", ",").split(",") if part.strip()]


async def get_job(session: AsyncSession, job_id: UUID) -> ImportJob | None:
    result = await session.execute(select(ImportJob).where(ImportJob.id == job_id).options(*_job_options()))
    return result.scalar_one_or_none()


async def create_upload_job(session: AsyncSession, filename: str, content: bytes) -> ImportJob:
    rows = _parse_rows(filename, content)
    job = ImportJob(
        filename=filename,
        source_type=_source_type(filename),
        status=ImportJobStatus.UPLOADED,
        raw_rows=rows,
        preview_rows=[],
        created_question_ids=[],
    )
    session.add(job)
    await session.flush()
    session.add(ImportSummary(job_id=job.id, total_rows=len(rows), valid_rows=0, invalid_rows=0, committed_rows=0, skipped_rows=0))
    await session.commit()
    return await get_job(session, job.id) or job


async def _get_or_create_category(session: AsyncSession, name: str) -> QuestionCategory:
    result = await session.execute(select(QuestionCategory).where(QuestionCategory.name == name))
    category = result.scalar_one_or_none()
    if category is None:
        category = QuestionCategory(name=name, description=f"{name} questions")
        session.add(category)
        await session.flush()
    return category


async def _find_course(session: AsyncSession, title: str | None) -> Course | None:
    if not title:
        return None
    result = await session.execute(select(Course).where(Course.title == title))
    return result.scalar_one_or_none()


async def _find_chapter(session: AsyncSession, title: str | None) -> Chapter | None:
    if not title:
        return None
    result = await session.execute(select(Chapter).where(Chapter.title == title))
    return result.scalar_one_or_none()


async def _find_lesson(session: AsyncSession, title: str | None) -> Lesson | None:
    if not title:
        return None
    result = await session.execute(select(Lesson).where(Lesson.title == title))
    return result.scalar_one_or_none()


async def _find_concepts(session: AsyncSession, names: list[str]) -> list[Concept]:
    if not names:
        return []
    result = await session.execute(select(Concept).where(Concept.name.in_(names)))
    return list(result.scalars().all())


async def _normalize_row(session: AsyncSession, row: dict, row_number: int) -> tuple[QuestionCreate | None, list[ImportError]]:
    errors: list[ImportError] = []
    title = _get(row, "question", "title", "prompt")
    if not title:
        errors.append(ImportError(row_number=row_number, field_name="question", message="Question text is required"))

    raw_type = (_get(row, "question_type", "type") or "MCQ").upper().replace(" ", "_").replace("-", "_")
    if raw_type == "SHORT":
        raw_type = "SHORT_ANSWER"
    if raw_type not in QuestionType.__members__:
        errors.append(ImportError(row_number=row_number, field_name="question_type", message=f"Unsupported question type: {raw_type}"))

    raw_difficulty = (_get(row, "difficulty") or "MEDIUM").upper()
    if raw_difficulty not in Difficulty.__members__:
        errors.append(ImportError(row_number=row_number, field_name="difficulty", message=f"Unsupported difficulty: {raw_difficulty}"))

    category = await _get_or_create_category(session, _get(row, "subject", "category") or "Chemistry")
    course = await _find_course(session, _get(row, "course"))
    chapter = await _find_chapter(session, _get(row, "chapter"))
    lesson = await _find_lesson(session, _get(row, "lesson"))
    concepts = await _find_concepts(session, _split(_get(row, "concept", "concepts")))

    choices: list[QuestionChoiceCreate] = []
    raw_choices = _split(_get(row, "choices", "options"))
    correct_answers = set(_split(_get(row, "correct_answer", "answer")))
    for index, choice_text in enumerate(raw_choices):
        choices.append(QuestionChoiceCreate(choice_text=choice_text, is_correct=choice_text in correct_answers, position=index))

    if errors or title is None:
        return None, errors

    payload = QuestionCreate(
        category_id=category.id,
        title=title,
        question_type=QuestionType[raw_type],
        difficulty=Difficulty[raw_difficulty],
        explanation=_get(row, "explanation"),
        correct_answer=_get(row, "correct_answer", "answer"),
        source=_get(row, "source"),
        bloom_level=_get(row, "bloom_level", "bloom"),
        thinking_skill=_get(row, "thinking_skill", "skill"),
        estimated_time_seconds=int(_get(row, "estimated_time_seconds", "estimated_time") or 60),
        common_mistakes=_split(_get(row, "common_mistakes", "mistakes")),
        keywords=_split(_get(row, "keywords")),
        course_id=course.id if course else None,
        chapter_id=chapter.id if chapter else None,
        lesson_id=lesson.id if lesson else None,
        concept_ids=[concept.id for concept in concepts],
        choices=choices,
        tags=_split(_get(row, "tags")),
    )
    return payload, []


async def preview_job(session: AsyncSession, job_id: UUID) -> ImportJob | None:
    job = await get_job(session, job_id)
    if job is None:
        return None

    for error in list(job.errors):
        await session.delete(error)
    await session.flush()

    preview_rows = []
    valid_rows = 0
    invalid_rows = 0
    for index, row in enumerate(job.raw_rows, start=1):
        payload, errors = await _normalize_row(session, row, index)
        if errors:
            invalid_rows += 1
            for error in errors:
                error.job_id = job.id
                session.add(error)
        else:
            valid_rows += 1
            preview_rows.append(payload.model_dump(mode="json") if payload else {})

    job.preview_rows = preview_rows
    job.status = ImportJobStatus.PREVIEWED if invalid_rows == 0 else ImportJobStatus.FAILED
    if job.summary is None:
        job.summary = ImportSummary(job_id=job.id)
    job.summary.total_rows = len(job.raw_rows)
    job.summary.valid_rows = valid_rows
    job.summary.invalid_rows = invalid_rows
    await session.commit()
    return await get_job(session, job_id)


async def commit_job(session: AsyncSession, job_id: UUID) -> ImportJob | None:
    job = await preview_job(session, job_id)
    if job is None or (job.summary and job.summary.invalid_rows > 0):
        return job

    created_question_ids: list[str] = []
    for row in job.preview_rows:
        payload = QuestionCreate(**row)
        question = await create_question(session, payload)
        created_question_ids.append(str(question.id))

    job = await get_job(session, job_id)
    if job is None:
        return None
    job.created_question_ids = created_question_ids
    job.status = ImportJobStatus.COMMITTED
    if job.summary is not None:
        job.summary.committed_rows = len(created_question_ids)
    await session.commit()
    return await get_job(session, job_id)


async def list_jobs(session: AsyncSession) -> list[ImportJob]:
    result = await session.execute(
        select(ImportJob)
        .options(*_job_options())
        .order_by(ImportJob.created_at.desc())
        .limit(50)
    )
    return list(result.scalars().all())
